"""Helpers for building Excel practice files (specs/practical-upgrade/spec.md §3).

Each lesson definition in defs/ describes its workbooks with `XlsxFile` objects. A workbook is
built twice: as a *starter* (data only, answer cells empty) and as a *finished* file (with the
real formulas). `expected` holds the correct answers, worked out in Python from the data, so the
checker can compare them with what Excel/LibreOffice calculates.
"""

import re
from dataclasses import dataclass, field
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Functions newer than Excel 2007 must be saved with an internal prefix, or Excel shows #NAME?
# until the formula is retyped.
FUTURE_FUNCTIONS = {
    'IFS': '_xlfn.', 'SWITCH': '_xlfn.', 'MAXIFS': '_xlfn.', 'MINIFS': '_xlfn.',
    'CONCAT': '_xlfn.', 'TEXTJOIN': '_xlfn.', 'XLOOKUP': '_xlfn.', 'XMATCH': '_xlfn.',
    'IFNA': '_xlfn.', 'DAYS': '_xlfn.', 'RANK.EQ': '_xlfn.', 'RANK.AVG': '_xlfn.',
    'UNIQUE': '_xlfn.', 'SEQUENCE': '_xlfn.', 'SORTBY': '_xlfn.',
    'SORT': '_xlfn._xlws.', 'FILTER': '_xlfn._xlws.',
}
# Functions that LibreOffice can only calculate from version 24.8.
NEEDS_NEW_LIBREOFFICE = {'XLOOKUP', 'XMATCH', 'SORT', 'SORTBY', 'FILTER', 'UNIQUE', 'SEQUENCE'}

_NAME = r'(?<![A-Za-z0-9_.])({})(?=\()'


def add_prefixes(formula: str) -> str:
    """Add Excel's internal prefix to every newer function in a formula."""
    for name in sorted(FUTURE_FUNCTIONS, key=len, reverse=True):
        formula = re.sub(_NAME.format(re.escape(name)), FUTURE_FUNCTIONS[name] + name, formula)
    return formula


def functions_used(formula: str) -> set:
    return {m.group(1).upper() for m in re.finditer(r'([A-Za-z][A-Za-z0-9.]*)\(', formula)}


def unprefixed_future_functions(formula: str) -> set:
    """Newer functions in a saved formula that are missing their prefix."""
    bad = set()
    for name, prefix in FUTURE_FUNCTIONS.items():
        for m in re.finditer(r'([A-Za-z0-9_.]*)(?<![A-Za-z0-9_]){}\('.format(re.escape(name)), formula):
            if not m.group(1).endswith(prefix):
                bad.add(name)
    return bad


@dataclass
class XlsxFile:
    """One practice workbook. `build(wb, finished)` fills the workbook."""
    name: str                                   # e.g. 'weekly-sales' → weekly-sales-starter.xlsx
    build: Callable[[Workbook, bool], None]
    expected: dict = field(default_factory=dict)  # {'Sheet!B8' or 'B8': value} in the finished file
    # Which files get automatic column widths: 'both', 'finished' (starter keeps Excel's default
    # widths, e.g. when the lesson teaches widening columns) or 'none' (build() sets widths).
    fit: str = 'both'


def new_workbook(sheet_title: str = 'Sheet1') -> Workbook:
    wb = Workbook()
    wb.active.title = sheet_title
    return wb


def put_rows(ws, rows, start_row: int = 1, start_col: int = 1, bold_first: bool = True):
    """Write a list of rows. Strings starting with '=' are formulas (prefixes added)."""
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is None:
                continue
            if isinstance(value, str) and value.startswith('='):
                value = add_prefixes(value)
            cell = ws.cell(start_row + r, start_col + c, value)
            if bold_first and r == 0:
                cell.font = Font(bold=True)


def formula(ws, ref: str, text: str):
    ws[ref] = add_prefixes(text)


def bold(ws, *refs):
    for ref in refs:
        ws[ref].font = Font(bold=True)


def fit_columns(ws, minimum: int = 8, maximum: int = 40):
    """Make every used column wide enough to read its longest value.

    Cells inside a merged range (e.g. a title across A1:C1) are skipped, so a long title does not
    make its first column too wide.
    """
    merged = set()
    for rng in ws.merged_cells.ranges:
        for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row,
                                min_col=rng.min_col, max_col=rng.max_col):
            merged.update(c.coordinate for c in row)
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or cell.coordinate in merged:
                continue
            text = str(cell.value)
            length = 12 if text.startswith('=') else len(text)
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(maximum, width + 3))
