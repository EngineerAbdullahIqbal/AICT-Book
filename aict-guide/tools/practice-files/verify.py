"""Check every practice file (specs/practical-upgrade/spec.md §3, PU-02).

For each lesson definition in defs/:
  1. the starter and finished files exist, and the lesson page links to both;
  2. the starter file has no answers in the answer cells;
  3. every newer function in the finished file has Excel's internal prefix;
  4. LibreOffice recalculates the finished file, and every answer cell equals the expected value
     worked out in Python.
Exit code 1 if anything fails.

Usage (from aict-guide/):  python3 tools/practice-files/verify.py [lesson-folder ...]
"""

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from build import load_defs, output_dir  # noqa: E402
from xlsxkit import NEEDS_NEW_LIBREOFFICE, functions_used, unprefixed_future_functions  # noqa: E402

DOCS = HERE.parent.parent / 'docs'
RECALC_DIR = HERE / 'recalc-work'


def libreoffice():
    """The newest LibreOffice available, and its version as a tuple."""
    best = None
    for exe in ('/snap/bin/libreoffice', 'soffice', 'libreoffice'):
        path = shutil.which(exe) or (exe if Path(exe).exists() else None)
        if not path:
            continue
        out = subprocess.run([path, '--version'], capture_output=True, text=True, timeout=60).stdout
        m = re.search(r'(\d+)\.(\d+)', out)
        if m:
            version = (int(m.group(1)), int(m.group(2)))
            if best is None or version > best[1]:
                best = (path, version)
    return best


def split_ref(ref, wb):
    if '!' in ref:
        sheet, cell = ref.split('!')
        return wb[sheet], cell
    return wb.worksheets[0], ref


def same(actual, expected):
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return abs(actual - expected) < 1e-9
    return actual == expected


def check_lesson(lesson, lo, pending):
    problems = []
    page = next(DOCS.glob(f'{lesson.MODULE}/{lesson.LESSON_FILE}'), None)
    page_text = page.read_text() if page else ''
    if not page:
        problems.append(f'lesson page {lesson.MODULE}/{lesson.LESSON_FILE} not found')
    out = output_dir(lesson)
    for f in lesson.FILES:
        starter = out / f'{f.name}-starter.xlsx'
        finished = out / f'{f.name}-finished.xlsx'
        for path in (starter, finished):
            if not path.exists():
                problems.append(f'missing file {path.name}')
            elif f'/files/{lesson.APP}/{lesson.FOLDER}/{path.name}' not in page_text:
                problems.append(f'lesson page does not link {path.name}')
        if not (starter.exists() and finished.exists()):
            continue
        # 2. starter has no answers
        swb = load_workbook(starter)
        for ref in f.expected:
            ws, cell = split_ref(ref, swb)
            if ws[cell].value not in (None, ''):
                problems.append(f'{starter.name}: answer cell {ref} is not empty')
        # 3. prefixes
        fwb = load_workbook(finished)
        used = set()
        new_function_cells = set()   # 'Sheet!A1' of every cell filled by a new-function formula
        for ws in fwb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    value = c.value
                    is_array = hasattr(value, 'text')          # ArrayFormula (SORT, FILTER …)
                    text = value.text if is_array else value
                    if not (isinstance(text, str) and text.startswith('=')):
                        continue
                    names = functions_used(text.replace('_xlfn._xlws.', '').replace('_xlfn.', ''))
                    used |= names
                    missing = unprefixed_future_functions(text)
                    if missing:
                        problems.append(f'{finished.name}!{c.coordinate}: no prefix for {sorted(missing)}')
                    if names & NEEDS_NEW_LIBREOFFICE:
                        if is_array and value.ref:
                            for out_row in ws[value.ref]:
                                for out_cell in (out_row if isinstance(out_row, tuple) else (out_row,)):
                                    new_function_cells.add(f'{ws.title}!{out_cell.coordinate}')
                        else:
                            new_function_cells.add(f'{ws.title}!{c.coordinate}')
        # 4. recalculate and compare. With LibreOffice older than 24.8, answer cells whose own
        # formula uses XLOOKUP / SORT / FILTER / UNIQUE … cannot be calculated: they are listed as
        # pending (not checked yet); every other answer cell is still checked.
        old_lo = lo[1] < (24, 8)
        pending_cells = set()
        if old_lo:
            for ref in f.expected:
                ws, cell = split_ref(ref, fwb)
                if f'{ws.title}!{cell}' in new_function_cells:
                    pending_cells.add(ref)
        if pending_cells:
            pending.append(f'{finished.name}: {len(pending_cells)} answer cell(s) use '
                           f'{sorted(used & NEEDS_NEW_LIBREOFFICE)}; need LibreOffice 24.8+ to check')
        # The snap LibreOffice cannot use /tmp, so work inside the project (git-ignored).
        RECALC_DIR.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=RECALC_DIR) as tmp:
            subprocess.run([lo[0], '--headless', '--calc', '--convert-to',
                            'xlsx:Calc MS Excel 2007 XML', '--outdir', tmp, str(finished)],
                           capture_output=True, timeout=180)
            recalculated = Path(tmp) / finished.name
            if not recalculated.exists():
                problems.append(f'{finished.name}: LibreOffice could not recalculate it')
                continue
            rwb = load_workbook(recalculated, data_only=True)
            for ref, want in f.expected.items():
                if ref in pending_cells:
                    continue
                ws, cell = split_ref(ref, rwb)
                got = ws[cell].value
                if not same(got, want):
                    problems.append(f'{finished.name}!{ref}: expected {want!r}, got {got!r}')
    return problems


def main(only):
    lo = libreoffice()
    if not lo:
        print('LibreOffice not found')
        return 1
    print(f'Using {lo[0]} (version {lo[1][0]}.{lo[1][1]})')
    failed = 0
    for lesson in load_defs():
        if only and lesson.FOLDER not in only:
            continue
        pending = []
        problems = check_lesson(lesson, lo, pending)
        status = 'FAIL' if problems else ('PEND' if pending else 'OK  ')
        print(f'{status} {lesson.FOLDER}')
        for p in problems:
            print('     -', p)
        for p in pending:
            print('     ~ pending:', p)
        failed += bool(problems or pending)
    # Pending cells also fail the run: files are published only when every answer is checked.
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main(set(sys.argv[1:])))
