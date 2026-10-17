"""Lesson 7.4: Formatting Cells: Font, Fill, Borders and Alignment.

Starter = Shop List / Mobile Stock as they are at the END of Lesson 7.3 (headings not bold, so
that Ctrl+B in step 6 turns bold on).
Finished = the end of Lesson 7.4: a new row 1 with the title merged and centred across A1:C1,
bold, size 16; headings in row 2 bold, light fill, centred; All Borders on the table.

No formulas. `expected` checks the last row of the table, which moves down one row when the
title row is inserted (empty in the starter). The formatting is checked by the asserts below.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlsxkit import XlsxFile, new_workbook
from defs.x07_02_entering_editing_saving_data import HEADINGS, put_list
from defs.x07_03_rows_and_columns import SHOP_END_73, MOBILE_END_73, TALL_ROW

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'formatting-cells.mdx'
FOLDER = 'formatting-cells'

SHOP_TITLE = 'Hamza General Store'
MOBILE_TITLE = 'Imran Mobile Centre'
LIGHT_BLUE = 'FFDDEBF7'   # a light fill colour (step 7: "select a light colour")
LIGHT_GREEN = 'FFE2EFDA'
THIN = Side(style='thin')
ALL_BORDERS = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_formatted(ws, title, items, fill_colour):
    ws['A1'] = title
    for c, heading in enumerate(HEADINGS, start=1):
        ws.cell(2, c, heading)
    for r, row in enumerate(items, start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    # Title: Merge & Center, bold, size 16 (steps 1-5).
    ws.merge_cells('A1:C1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = TALL_ROW
    # The heading row (made taller in Lesson 7.3, now row 2): bold, fill, centre (steps 6-8).
    ws.row_dimensions[2].height = TALL_ROW
    for c in range(1, 4):
        cell = ws.cell(2, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type='solid', fgColor=fill_colour)
        cell.alignment = Alignment(horizontal='center')
    # All Borders from A2 to the last price (steps 10-11).
    for row in ws.iter_rows(min_row=2, max_row=2 + len(items), min_col=1, max_col=3):
        for cell in row:
            cell.border = ALL_BORDERS


def build_shop_list(wb, finished):
    ws = wb.active
    if finished:
        build_formatted(ws, SHOP_TITLE, SHOP_END_73, LIGHT_BLUE)
    else:
        put_list(ws, SHOP_END_73)
        ws.row_dimensions[1].height = TALL_ROW


def build_mobile_stock(wb, finished):
    ws = wb.active
    if finished:
        build_formatted(ws, MOBILE_TITLE, MOBILE_END_73, LIGHT_GREEN)
    else:
        put_list(ws, MOBILE_END_73)
        ws.row_dimensions[1].height = TALL_ROW


FILES = [
    XlsxFile('shop-list', build_shop_list, {'A9': 'Eggs', 'B9': 12, 'C9': 30}),
    XlsxFile('mobile-stock', build_mobile_stock, {'A8': 'Memory card', 'B8': 3, 'C8': 1200}),
]


def _check(build, title, last_row):
    wb = new_workbook()
    build(wb, False)
    assert not any(c.font.b for row in wb.active.iter_rows() for c in row)  # Ctrl+B turns bold on
    wb = new_workbook()
    build(wb, True)
    ws = wb.active
    assert ws['A1'].value == title and 'A1:C1' in [str(r) for r in ws.merged_cells.ranges]
    assert ws['A1'].font.b and ws['A1'].font.sz == 16 and ws['A1'].alignment.horizontal == 'center'
    assert [ws.cell(2, c).value for c in range(1, 4)] == list(HEADINGS)
    for c in range(1, 4):
        cell = ws.cell(2, c)
        assert cell.font.b and cell.fill.fill_type == 'solid' and cell.alignment.horizontal == 'center'
    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=3):
        for cell in row:
            assert all(getattr(cell.border, s).style == 'thin' for s in ('left', 'right', 'top', 'bottom'))
    assert ws.max_row == last_row


# The lesson states these results; keep them in step with the page.
_check(build_shop_list, SHOP_TITLE, 9)        # table A2 to C9
_check(build_mobile_stock, MOBILE_TITLE, 8)   # Practice: "from A2 to the last price in C8"
