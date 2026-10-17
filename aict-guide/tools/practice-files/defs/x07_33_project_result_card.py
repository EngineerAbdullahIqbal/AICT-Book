"""Lesson 7.33: Project: Result Card With Grades.

Two result sheets with the same layout:
  row 1 headings: A Roll No, B Name, C English, D Urdu, E Maths, F Science, G Computer, H Total,
  I Percentage, J Grade, K Result, L Position; students from row 2;
  one empty row, then Passed and Failed (COUNTIF on Result);
  one empty row, then the result card: Roll No (typed) and six VLOOKUP rows
  (Name, Total, Percentage, Grade, Result, Position).
Total = SUM; Percentage = Total / full marks * 100; Grade = IFS on the percentage with the bands
of Lesson 7.16 (80 A, 70 B, 60 C, 40 D, else F); Result = IF(AND(every subject >= pass mark));
Position = RANK.EQ(Total, $H$2:$H$last).

Starter: roll numbers, names, marks, the labels and the roll number to look up. Finished: the
formulas and a "Total Marks" column chart of the names and totals. The finished files show the
end of each task (Step by Step: roll number 107; Practice: roll number 203).
"""

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef

from xlsxkit import XlsxFile, new_workbook, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-result-card.mdx'
FOLDER = 'project-result-card'

HEADINGS = ('Roll No', 'Name', 'English', 'Urdu', 'Maths', 'Science', 'Computer', 'Total',
            'Percentage', 'Grade', 'Result', 'Position')
CARD = ('Name', 'Total', 'Percentage', 'Grade', 'Result', 'Position')
CARD_COLUMNS = (2, 8, 9, 10, 11, 12)          # column numbers inside A:L

# Step by Step: a girls' and boys' high school in Bahawalpur, Grade 9 term test, each subject
# out of 100, pass mark 40.
SCHOOL = dict(full=500, pass_mark=40, look_start=104, look_end=107, students=[
    (101, 'Ayesha Siddiqui', 82, 78, 91, 85, 88),
    (102, 'Hassan Raza', 65, 70, 58, 62, 75),
    (103, 'Mahnoor Fatima', 74, 81, 69, 77, 80),
    (104, 'Ali Haider', 55, 60, 35, 48, 66),
    (105, 'Iqra Javed', 90, 85, 88, 92, 95),
    (106, 'Usman Ghani', 45, 52, 41, 38, 50),
    (107, 'Sidra Malik', 68, 72, 75, 70, 79),
    (108, 'Zeeshan Akhtar', 30, 42, 28, 35, 40),
])
# Practice: a tuition centre's monthly test, each subject out of 50, pass mark 20.
TUITION = dict(full=250, pass_mark=20, look_start=203, look_end=203, students=[
    (201, 'Hamza Iqbal', 42, 38, 45, 40, 44),
    (202, 'Areeba Khan', 35, 40, 30, 38, 37),
    (203, 'Fahad Mirza', 25, 30, 18, 28, 32),
    (204, 'Noor Fatima', 40, 44, 39, 41, 45),
    (205, 'Saad Anwar', 30, 33, 35, 29, 31),
    (206, 'Laiba Aslam', 20, 24, 22, 21, 26),
])


def rows_for(sheet):
    last = 1 + len(sheet['students'])
    passed, failed = last + 2, last + 3
    roll = last + 5
    return last, passed, failed, roll


def make_builder(sheet):
    last, passed_row, failed_row, roll_row = rows_for(sheet)
    full, pm = sheet['full'], sheet['pass_mark']

    def build(wb, finished):
        ws = wb.active
        put_rows(ws, [HEADINGS, *sheet['students']])
        ws[f'A{passed_row}'], ws[f'A{failed_row}'] = 'Passed', 'Failed'
        ws[f'A{roll_row}'] = 'Roll No'
        ws[f'B{roll_row}'] = sheet['look_end'] if finished else sheet['look_start']
        for k, label in enumerate(CARD, start=1):
            ws[f'A{roll_row + k}'] = label
        bold(ws, f'A{passed_row}', f'A{failed_row}',
             *(f'A{roll_row + k}' for k in range(len(CARD) + 1)))
        if not finished:
            return
        for r in range(2, last + 1):
            formula(ws, f'H{r}', f'=SUM(C{r}:G{r})')
            formula(ws, f'I{r}', f'=H{r}/{full}*100')
            formula(ws, f'J{r}', f'=IFS(I{r}>=80,"A",I{r}>=70,"B",I{r}>=60,"C",I{r}>=40,"D",TRUE,"F")')
            tests = ','.join(f'{c}{r}>={pm}' for c in 'CDEFG')
            formula(ws, f'K{r}', f'=IF(AND({tests}),"Pass","Fail")')
            formula(ws, f'L{r}', f'=RANK.EQ(H{r},$H$2:$H${last})')
        formula(ws, f'B{passed_row}', f'=COUNTIF(K2:K{last},"Pass")')
        formula(ws, f'B{failed_row}', f'=COUNTIF(K2:K{last},"Fail")')
        for k, col_num in enumerate(CARD_COLUMNS, start=1):
            formula(ws, f'B{roll_row + k}', f'=VLOOKUP(B{roll_row},A2:L{last},{col_num},FALSE)')
        chart = BarChart()
        chart.type = 'col'
        chart.grouping = 'clustered'
        chart.title = 'Total Marks'
        chart.add_data(Reference(ws, min_col=8, min_row=1, max_row=last), titles_from_data=True)
        # The names are text, so the categories are a text reference (strRef).
        chart.series[0].cat = AxDataSource(strRef=StrRef(f=f"'{ws.title}'!$B$2:$B${last}"))
        chart.legend = None
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        ws.add_chart(chart, 'N2')
    return build


def grade(pct):
    for limit, g in ((80, 'A'), (70, 'B'), (60, 'C'), (40, 'D')):
        if pct >= limit:
            return g
    return 'F'


def results(sheet):
    """Per student: (roll, name, total, percentage, grade, result, position)."""
    out = []
    totals = [sum(s[2:]) for s in sheet['students']]
    for s, total in zip(sheet['students'], totals):
        pct = total / sheet['full'] * 100
        result = 'Pass' if all(m >= sheet['pass_mark'] for m in s[2:]) else 'Fail'
        position = 1 + sum(t > total for t in totals)       # RANK.EQ, largest first
        out.append((s[0], s[1], total, pct, grade(pct), result, position))
    return out


def expected_for(sheet):
    last, passed_row, failed_row, roll_row = rows_for(sheet)
    res = results(sheet)
    exp = {}
    for r, (_, _, total, pct, g, result, pos) in enumerate(res, start=2):
        exp.update({f'H{r}': total, f'I{r}': pct, f'J{r}': g, f'K{r}': result, f'L{r}': pos})
    exp[f'B{passed_row}'] = sum(x[5] == 'Pass' for x in res)
    exp[f'B{failed_row}'] = sum(x[5] == 'Fail' for x in res)
    card = next(x for x in res if x[0] == sheet['look_end'])
    for k, value in enumerate((card[1], card[2], card[3], card[4], card[5], card[6]), start=1):
        exp[f'B{roll_row + k}'] = value
    return exp


school = expected_for(SCHOOL)
tuition = expected_for(TUITION)

FILES = [
    XlsxFile('school-result', make_builder(SCHOOL), school),
    XlsxFile('tuition-result', make_builder(TUITION), tuition),
]

# The lesson states these results; keep them in step with the page.
assert rows_for(SCHOOL) == (9, 11, 12, 14) and rows_for(TUITION) == (7, 9, 10, 12)
assert [(t, round(p, 6), g, r, pos) for _, _, t, p, g, r, pos in results(SCHOOL)] == [
    (424, 84.8, 'A', 'Pass', 2), (330, 66, 'C', 'Pass', 5), (381, 76.2, 'B', 'Pass', 3),
    (264, 52.8, 'D', 'Fail', 6), (450, 90, 'A', 'Pass', 1), (226, 45.2, 'D', 'Fail', 7),
    (364, 72.8, 'B', 'Pass', 4), (175, 35, 'F', 'Fail', 8)]
assert school['B11'] == 5 and school['B12'] == 3
card_104 = next(x for x in results(SCHOOL) if x[0] == 104)
assert (card_104[1], card_104[2], round(card_104[3], 6), *card_104[4:]) == (
    'Ali Haider', 264, 52.8, 'D', 'Fail', 6)
assert (school['B15'], school['B16'], round(school['B17'], 6), school['B18'], school['B19'],
        school['B20']) == ('Sidra Malik', 364, 72.8, 'B', 'Pass', 4)

assert [(t, round(p, 6), g, r, pos) for _, _, t, p, g, r, pos in results(TUITION)] == [
    (209, 83.6, 'A', 'Pass', 1), (180, 72, 'B', 'Pass', 3), (133, 53.2, 'D', 'Fail', 5),
    (209, 83.6, 'A', 'Pass', 1), (158, 63.2, 'C', 'Pass', 4), (113, 45.2, 'D', 'Pass', 6)]
assert tuition['B9'] == 5 and tuition['B10'] == 1
assert (tuition['B13'], tuition['B14'], round(tuition['B15'], 6), tuition['B16'],
        tuition['B17'], tuition['B18']) == ('Fahad Mirza', 133, 53.2, 'D', 'Fail', 5)

# The finished files have one chart; the starters have none and leave the answers empty.
for _sheet in (SCHOOL, TUITION):
    for _finished in (False, True):
        _wb = new_workbook()
        make_builder(_sheet)(_wb, _finished)
        assert len(_wb.active._charts) == (1 if _finished else 0)
