"""Lesson 7.34: Project: Tuition Fee Record and Unpaid List.

Two fee records with the same layout:
  A1 'Due date', B1 the due date (one cell, used with $B$1);
  row 3 headings: A Student, B Group, C Fee (Rs), D Paid (Rs), E Paid On, F Balance (Rs),
  G Status, H Days Late; students from row 4;
  one empty row, then the summary: Unpaid students (COUNTIF), unpaid balance per group (SUMIF)
  and Total unpaid (SUM).
Balance = C-D; Status = IF(F>0,"Unpaid","Paid"); Days Late = IF(E>$B$1,E-$B$1,0).

Dates are real Excel dates in Short Date (built-in format 14), like Lesson 7.24. Answer cells
stay General, so days show as numbers.

Starter: data only. Finished: the end of the task: the formulas, a "Text that Contains Unpaid"
rule (red) on the Status cells, and the filter on Status showing only "Unpaid" (openpyxl stores
the filter but does not hide rows, so build() hides them, as Excel does). The Practice finished
file shows Faizan's payment changed to 4000 on 8 December.
"""

from datetime import datetime

from openpyxl.formatting.rule import Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle

from xlsxkit import XlsxFile, new_workbook, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-fee-record.mdx'
FOLDER = 'project-fee-record'

SHORT_DATE = 'mm-dd-yy'          # openpyxl's name for Excel's built-in format 14 (Short Date)
HEADINGS = ('Student', 'Group', 'Fee (Rs)', 'Paid (Rs)', 'Paid On', 'Balance (Rs)', 'Status',
            'Days Late')
UNPAID = 'Unpaid'
FIRST = 4
RED = DifferentialStyle(fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE',
                                         fill_type='solid'), font=Font(color='9C0006'))


def dt(y, m, d):
    return datetime(y, m, d)


# Step by Step: a tuition centre in Lahore, November fee due on 10 November 2026.
CENTRE = dict(due=dt(2026, 11, 10), groups=['Grade 8', 'Grade 9', 'Grade 10'], students=[
    ('Ahmed Raza', 'Grade 8', 2500, 2500, dt(2026, 11, 8)),
    ('Zainab Ali', 'Grade 8', 2500, 1500, dt(2026, 11, 12)),
    ('Hamid Nawaz', 'Grade 9', 3000, 3000, dt(2026, 11, 10)),
    ('Sana Tariq', 'Grade 9', 3000, 2000, dt(2026, 11, 15)),
    ('Omer Farooq', 'Grade 9', 3000, 3000, dt(2026, 11, 14)),
    ('Hina Akram', 'Grade 10', 3500, 3500, dt(2026, 11, 9)),
    ('Danish Iqbal', 'Grade 10', 3500, 1000, dt(2026, 11, 20)),
    ('Maham Shah', 'Grade 8', 2500, 2500, dt(2026, 11, 11)),
    ('Fahad Butt', 'Grade 10', 3500, 2000, dt(2026, 11, 18)),
    ('Amna Yousaf', 'Grade 9', 3000, 1500, dt(2026, 11, 25)),
])
# Practice: a computer course centre, December fee due on 5 December 2026.
COURSE = dict(due=dt(2026, 12, 5), groups=['Morning', 'Evening'], students=[
    ('Rizwan Ali', 'Morning', 4000, 4000, dt(2026, 12, 3)),
    ('Saima Noreen', 'Evening', 3500, 2000, dt(2026, 12, 7)),
    ('Jawad Khan', 'Morning', 4000, 2500, dt(2026, 12, 9)),
    ('Ayesha Rehman', 'Evening', 3500, 3500, dt(2026, 12, 5)),
    ('Tahir Mehmood', 'Morning', 4000, 4000, dt(2026, 12, 12)),
    ('Nimra Saleem', 'Evening', 3500, 1000, dt(2026, 12, 15)),
    ('Faizan Ahmed', 'Morning', 4000, 3000, dt(2026, 12, 6)),
    ('Komal Aziz', 'Evening', 3500, 3500, dt(2026, 12, 4)),
])
# Practice step: Faizan pays the other 1000 on 8 December, so D10 becomes 4000 and E10 8/12/2026.
COURSE_CHANGED = dict(COURSE, students=[
    (s[0], s[1], s[2], 4000, dt(2026, 12, 8)) if s[0] == 'Faizan Ahmed' else s
    for s in COURSE['students']])


def summary_labels(sheet):
    return ['Unpaid students', *(f'Unpaid {g} (Rs)' for g in sheet['groups']), 'Total unpaid (Rs)']


def rows_for(sheet):
    last = FIRST + len(sheet['students']) - 1
    return last, last + 2                        # last student row, first summary row


def make_builder(start, end):
    last, first_sum = rows_for(start)

    def build(wb, finished):
        ws = wb.active
        sheet = end if finished else start
        ws['A1'], ws['B1'] = 'Due date', sheet['due']
        ws['B1'].number_format = SHORT_DATE
        bold(ws, 'A1')
        put_rows(ws, [HEADINGS, *sheet['students']], start_row=FIRST - 1)
        for r in range(FIRST, last + 1):
            ws[f'E{r}'].number_format = SHORT_DATE
        labels = summary_labels(sheet)
        for k, label in enumerate(labels):
            ws[f'A{first_sum + k}'] = label
            bold(ws, f'A{first_sum + k}')
        if not finished:
            return
        for r in range(FIRST, last + 1):
            formula(ws, f'F{r}', f'=C{r}-D{r}')
            formula(ws, f'G{r}', f'=IF(F{r}>0,"{UNPAID}","Paid")')
            formula(ws, f'H{r}', f'=IF(E{r}>$B$1,E{r}-$B$1,0)')
        formula(ws, f'B{first_sum}', f'=COUNTIF(G{FIRST}:G{last},"{UNPAID}")')
        for k, g in enumerate(sheet['groups'], start=1):
            formula(ws, f'B{first_sum + k}',
                    f'=SUMIF(B{FIRST}:B{last},"{g}",F{FIRST}:F{last})')
        formula(ws, f'B{first_sum + len(labels) - 1}', f'=SUM(F{FIRST}:F{last})')
        ws.conditional_formatting.add(f'G{FIRST}:G{last}', Rule(
            type='containsText', operator='containsText', text=UNPAID,
            formula=[f'NOT(ISERROR(SEARCH("{UNPAID}",G{FIRST})))'], dxf=RED))
        ws.auto_filter.ref = f'A{FIRST - 1}:H{last}'
        ws.auto_filter.add_filter_column(6, [UNPAID])            # column G (Status)
        for r, s in enumerate(sheet['students'], start=FIRST):
            if s[2] - s[3] <= 0:
                ws.row_dimensions[r].hidden = True
    return build


def expected_for(sheet):
    last, first_sum = rows_for(sheet)
    exp = {}
    for r, (_, _, fee, paid, paid_on) in enumerate(sheet['students'], start=FIRST):
        late = (paid_on - sheet['due']).days
        exp[f'F{r}'] = fee - paid
        exp[f'G{r}'] = UNPAID if fee - paid > 0 else 'Paid'
        exp[f'H{r}'] = late if late > 0 else 0
    exp[f'B{first_sum}'] = sum(s[2] > s[3] for s in sheet['students'])
    for k, g in enumerate(sheet['groups'], start=1):
        exp[f'B{first_sum + k}'] = sum(s[2] - s[3] for s in sheet['students'] if s[1] == g)
    exp[f'B{first_sum + len(sheet["groups"]) + 1}'] = sum(s[2] - s[3] for s in sheet['students'])
    return exp


centre = expected_for(CENTRE)
course_before = expected_for(COURSE)
course = expected_for(COURSE_CHANGED)

FILES = [
    XlsxFile('tuition-fees', make_builder(CENTRE, CENTRE), centre),
    XlsxFile('course-fees', make_builder(COURSE, COURSE_CHANGED), course),
]


# The lesson states these results; keep them in step with the page.
def col(exp, c, first, last):
    return [exp[f'{c}{r}'] for r in range(first, last + 1)]


assert rows_for(CENTRE) == (13, 15) and rows_for(COURSE) == (11, 13)
assert col(centre, 'F', 4, 13) == [0, 1000, 0, 1000, 0, 0, 2500, 0, 1500, 1500]
assert col(centre, 'G', 4, 13) == ['Paid', UNPAID, 'Paid', UNPAID, 'Paid', 'Paid', UNPAID,
                                   'Paid', UNPAID, UNPAID]
assert col(centre, 'H', 4, 13) == [0, 2, 0, 5, 4, 0, 10, 1, 8, 15]
assert col(centre, 'B', 15, 19) == [5, 1000, 2500, 4000, 7500]

assert col(course_before, 'F', 4, 11) == [0, 1500, 1500, 0, 0, 2500, 1000, 0]
assert col(course_before, 'H', 4, 11) == [0, 2, 4, 0, 7, 10, 1, 0]
assert col(course_before, 'B', 13, 16) == [4, 2500, 4000, 6500]
assert col(course, 'F', 4, 11) == [0, 1500, 1500, 0, 0, 2500, 0, 0]
assert col(course, 'G', 4, 11) == ['Paid', UNPAID, UNPAID, 'Paid', 'Paid', UNPAID, 'Paid',
                                   'Paid']
assert col(course, 'H', 4, 11) == [0, 2, 4, 0, 7, 10, 3, 0]
assert col(course, 'B', 13, 16) == [3, 1500, 4000, 5500]

# The finished files: the filter shows only the Unpaid rows; the starters have no filter.
for _start, _end, _shown in ((CENTRE, CENTRE, [5, 7, 10, 12, 13]),
                             (COURSE, COURSE_CHANGED, [5, 6, 9])):
    _wb = new_workbook()
    make_builder(_start, _end)(_wb, True)
    _ws = _wb.active
    _last = rows_for(_start)[0]
    assert _ws.auto_filter.ref == f'A3:H{_last}'
    assert [r for r in range(FIRST, _last + 1) if not _ws.row_dimensions[r].hidden] == _shown
    assert len(_ws.conditional_formatting) == 1
    _wb = new_workbook()
    make_builder(_start, _end)(_wb, False)
    assert _wb.active.auto_filter.ref is None and len(_wb.active.conditional_formatting) == 0
