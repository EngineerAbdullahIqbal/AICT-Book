"""Lesson 7.35: Project: Monthly Family Budget.

Two budgets with the same layout:
  row 1 headings Income | Amount (Rs); rows 2-3 the two incomes; A4 Total income, B4 =SUM(B2:B3);
  row 6 headings: A Expense, B Planned (Rs), C Actual (Rs), D Difference (Rs),
  E Share of Income (%); the eight expenses in rows 7-14; row 15 Total;
  A17 Money left (Rs), B17 =B4-C15.
Difference = Actual - Planned (a plus number means more was spent than planned);
Share of Income = ROUND(Actual / $B$4 * 100, 1). Row 15: B15 and C15 are SUMs; D15 and E15 are
the row 7 formulas filled down to row 15.

Starter: incomes, expenses, planned and actual amounts (data only). Finished: the formulas, the
Number format with the thousands separator on the money cells, a conditional formatting rule
(Greater Than 0, red) on D7:D15 for overspending, and a clustered column chart of Planned and
Actual for A6:C14. The Practice finished file shows the end of the task, with Transport changed
to 4900.
"""

import math

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from xlsxkit import XlsxFile, new_workbook, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-family-budget.mdx'
FOLDER = 'project-family-budget'

HEADINGS = ('Expense', 'Planned (Rs)', 'Actual (Rs)', 'Difference (Rs)', 'Share of Income (%)')
EXPENSES = ('House rent', 'Groceries', 'Electricity', 'Gas', 'School fees', 'Transport',
            'Mobile', 'Savings')
FIRST, LAST, TOTAL, LEFT = 7, 14, 15, 17
MONEY = '#,##0'              # Number, 0 decimal places, Use 1000 Separator
RED = dict(fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
           font=Font(color='9C0006'))
CF_RANGE = f'D{FIRST}:D{TOTAL}'
CHART_TITLE = 'Planned and Actual Spending'

# Step by Step: Tariq and Shabana's family in Faisalabad, October.
OCTOBER = dict(
    incomes=[('Salary', 78000), ('Home tailoring', 14000)],
    planned=[22000, 20000, 7500, 2500, 8500, 5000, 1500, 10000],
    actual=[22000, 21650, 8930, 2180, 8500, 5640, 1350, 10000],
)
# Practice: Rashid and Nargis's family in Peshawar. A missing rickshaw receipt later raises
# Transport (row 12) from 4300 to 4900.
PESHAWAR = dict(
    incomes=[('Shop income', 62000), ('Home tuition', 15000)],
    planned=[18000, 19000, 6500, 2800, 7000, 4500, 1200, 8000],
    actual=[18000, 18400, 7860, 3150, 7000, 4300, 1290, 8000],
)
TRANSPORT_ROW, TRANSPORT_NEW = 12, 4900


def peshawar_after_change():
    actual = list(PESHAWAR['actual'])
    actual[TRANSPORT_ROW - FIRST] = TRANSPORT_NEW
    return dict(PESHAWAR, actual=actual)


def add_budget_chart(ws):
    chart = BarChart()
    chart.type = 'col'
    chart.grouping = 'clustered'
    chart.title = CHART_TITLE
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=FIRST - 1, max_row=LAST),
                   titles_from_data=True)
    cats = StrRef(f=f"'{ws.title}'!$A${FIRST}:$A${LAST}")
    for s in chart.series:
        s.cat = AxDataSource(strRef=cats)
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.width, chart.height = 18, 9
    ws.add_chart(chart, 'G2')


def make_builder(budget):
    def build(wb, finished):
        ws = wb.active
        put_rows(ws, [('Income', 'Amount (Rs)'), *budget['incomes'], ('Total income',)])
        bold(ws, 'A4')
        rows = [(name, p, a) for name, p, a in zip(EXPENSES, budget['planned'], budget['actual'])]
        put_rows(ws, [HEADINGS, *rows], start_row=FIRST - 1)
        ws[f'A{TOTAL}'] = 'Total'
        ws[f'A{LEFT}'] = 'Money left (Rs)'
        bold(ws, f'A{TOTAL}', f'A{LEFT}')
        if not finished:
            return
        formula(ws, 'B4', '=SUM(B2:B3)')
        formula(ws, f'B{TOTAL}', f'=SUM(B{FIRST}:B{LAST})')
        formula(ws, f'C{TOTAL}', f'=SUM(C{FIRST}:C{LAST})')
        for r in range(FIRST, TOTAL + 1):
            formula(ws, f'D{r}', f'=C{r}-B{r}')
            formula(ws, f'E{r}', f'=ROUND(C{r}/$B$4*100,1)')
        formula(ws, f'B{LEFT}', f'=B4-C{TOTAL}')
        bold(ws, f'B{LEFT}')
        for ref in ('B2', 'B3', 'B4', f'B{LEFT}'):
            ws[ref].number_format = MONEY
        for r in range(FIRST, TOTAL + 1):
            for col in 'BCD':
                ws[f'{col}{r}'].number_format = MONEY
        ws.conditional_formatting.add(CF_RANGE, CellIsRule(operator='greaterThan',
                                                           formula=['0'], **RED))
        add_budget_chart(ws)
    return build


def excel_round(x, digits):
    """ROUND(x, digits) for the positive numbers here: halves round up, as in Excel."""
    v = x * 10 ** digits
    assert abs(v - math.floor(v) - 0.5) > 1e-6, 'avoid exact halves in practice data'
    return math.floor(v + 0.5) / 10 ** digits


def expected_for(budget):
    income = sum(v for _, v in budget['incomes'])
    exp = {'B4': income}
    for r, (p, a) in enumerate(zip(budget['planned'], budget['actual']), start=FIRST):
        exp[f'D{r}'] = a - p
        exp[f'E{r}'] = excel_round(a / income * 100, 1)
    planned, actual = sum(budget['planned']), sum(budget['actual'])
    exp[f'B{TOTAL}'] = planned
    exp[f'C{TOTAL}'] = actual
    exp[f'D{TOTAL}'] = actual - planned
    exp[f'E{TOTAL}'] = excel_round(actual / income * 100, 1)
    exp[f'B{LEFT}'] = income - actual
    return exp


def over_plan(exp):
    """Cells the Greater Than 0 rule colours red."""
    return [f'D{r}' for r in range(FIRST, TOTAL + 1) if exp[f'D{r}'] > 0]


october = expected_for(OCTOBER)
peshawar_before = expected_for(PESHAWAR)
peshawar = expected_for(peshawar_after_change())

FILES = [
    XlsxFile('october-budget', make_builder(OCTOBER), october),
    XlsxFile('peshawar-budget', make_builder(peshawar_after_change()), peshawar),
]


# The lesson states these results; keep them in step with the page.
def column(exp, col):
    return [exp[f'{col}{r}'] for r in range(FIRST, TOTAL + 1)]


assert october['B4'] == 92000
assert column(october, 'D') == [0, 1650, 1430, -320, 0, 640, -150, 0, 3250]
assert column(october, 'E') == [23.9, 23.5, 9.7, 2.4, 9.2, 6.1, 1.5, 10.9, 87.2]
assert (october['B15'], october['C15'], october['B17']) == (77000, 80250, 11750)
assert over_plan(october) == ['D8', 'D9', 'D12', 'D15']

assert peshawar_before['B4'] == 77000
assert column(peshawar_before, 'D') == [0, -600, 1360, 350, 0, -200, 90, 0, 1000]
assert (peshawar_before['B15'], peshawar_before['C15'], peshawar_before['B17']) == (
    67000, 68000, 9000)
assert peshawar_before['E15'] == 88.3
assert over_plan(peshawar_before) == ['D9', 'D10', 'D13', 'D15']
assert column(peshawar, 'D') == [0, -600, 1360, 350, 0, 400, 90, 0, 1600]
assert column(peshawar, 'E') == [23.4, 23.9, 10.2, 4.1, 9.1, 6.4, 1.7, 10.4, 89.1]
assert (peshawar['B15'], peshawar['C15'], peshawar['B17']) == (67000, 68600, 8400)
assert over_plan(peshawar) == ['D9', 'D10', 'D12', 'D13', 'D15']
# The chart's tallest Actual column (row 7 to 14).
assert max(zip(OCTOBER['actual'], EXPENSES)) == (22000, 'House rent')
assert sorted(zip(OCTOBER['actual'], EXPENSES))[-2] == (21650, 'Groceries')

# The finished files hold one rule (Greater Than 0 on D7:D15) and one chart; starters hold none.
for _budget in (OCTOBER, peshawar_after_change()):
    _wb = new_workbook()
    make_builder(_budget)(_wb, True)
    _ws = _wb.active
    assert [(str(cf.sqref), rule.type, rule.operator, list(rule.formula))
            for cf in _ws.conditional_formatting for rule in cf.rules] == [
        (CF_RANGE, 'cellIs', 'greaterThan', ['0'])]
    assert len(_ws._charts) == 1
    _wb = new_workbook()
    make_builder(_budget)(_wb, False)
    assert not list(_wb.active.conditional_formatting) and not _wb.active._charts
