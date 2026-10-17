"""Lesson 7.25: Excel 2021 Formulas: SORT, FILTER and UNIQUE.

In the finished files the dynamic array formulas are saved as openpyxl ArrayFormula objects over
the exact output range (writer brief), with Excel's prefixes, so that the results show in Excel
2021. Each `expected` value is worked out in Python from the data.

verify.py cannot recalculate these files until LibreOffice 24.8+ is installed.
"""

from openpyxl.worksheet.formula import ArrayFormula

from xlsxkit import XlsxFile, put_rows, add_prefixes, fit_columns

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'sort-filter-unique.mdx'
FOLDER = 'sort-filter-unique'

HEAD = ('Day', 'Product', 'Amount (Rs)')

# Step by Step: one week of sales at a small general store (made-up amounts).
SALES = [('Monday', 'Sugar', 1200), ('Monday', 'Tea', 850), ('Tuesday', 'Rice', 2400),
         ('Tuesday', 'Sugar', 1500), ('Wednesday', 'Cooking oil', 3100),
         ('Wednesday', 'Tea', 950), ('Thursday', 'Rice', 1800),
         ('Thursday', 'Cooking oil', 2700), ('Friday', 'Sugar', 900), ('Friday', 'Tea', 1300)]
OVER = 2000


def unique(values):
    """Distinct values in order of first appearance, like UNIQUE."""
    out = []
    for v in values:
        if v not in out:
            out.append(v)
    return out


def array_formula(ws, top_left, rows, cols, text):
    """Write a dynamic array formula over its exact output range."""
    from openpyxl.utils import get_column_letter, column_index_from_string
    from openpyxl.utils.cell import coordinate_from_string
    col, row = coordinate_from_string(top_left)
    last = get_column_letter(column_index_from_string(col) + cols - 1) + str(row + rows - 1)
    ws[top_left] = ArrayFormula(ref=f'{top_left}:{last}', text=add_prefixes(text))


def fit_widths(ws):
    """Column widths for both files (fit='none'): fit the typed data, then give each output
    column the width of the data column it copies, so the spilled results are readable."""
    fit_columns(ws)
    dims = ws.column_dimensions
    for out, src in (('E', 'B'), ('G', 'A'), ('H', 'B'), ('I', 'C'), ('K', 'A'), ('L', 'B'),
                     ('M', 'C')):
        dims[out].width = max(dims[out].width or 0, dims[src].width)


def put_expected(expected, top_left_col, first_row, rows):
    from openpyxl.utils import get_column_letter, column_index_from_string
    c0 = column_index_from_string(top_left_col)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            expected[f'{get_column_letter(c0 + j)}{first_row + i}'] = v


products = unique([p for _, p, _ in SALES])
by_amount = sorted(SALES, key=lambda s: s[2], reverse=True)
big_sales = [s for s in SALES if s[2] > OVER]


def build_shop_sales(wb, finished):
    ws = wb.active
    put_rows(ws, [HEAD, *SALES])
    put_rows(ws, [('Products',)], start_col=5)
    put_rows(ws, [HEAD], start_col=7)
    put_rows(ws, [HEAD], start_col=11)
    fit_widths(ws)
    if finished:
        array_formula(ws, 'E2', len(products), 1, '=UNIQUE(B2:B11)')
        array_formula(ws, 'G2', len(by_amount), 3, '=SORT(A2:C11,3,-1)')
        array_formula(ws, 'K2', len(big_sales), 3, '=FILTER(A2:C11,C2:C11>2000)')


sales_expected = {}
put_expected(sales_expected, 'E', 2, [(p,) for p in products])
put_expected(sales_expected, 'G', 2, by_amount)
put_expected(sales_expected, 'K', 2, big_sales)

# Practice: a tuition centre's test marks (made-up names and marks).
HEAD2 = ('Student', 'Grade', 'Marks')
MARKS = [('Hassan', 'Grade 8', 78), ('Noor', 'Grade 7', 91), ('Ali Raza', 'Grade 6', 64),
         ('Sana', 'Grade 8', 85), ('Usman', 'Grade 7', 72), ('Iqra', 'Grade 6', 59),
         ('Bilal', 'Grade 8', 88), ('Amna', 'Grade 7', 95)]
PICK = 'Grade 7'

grades = sorted(unique([c for _, c, _ in MARKS]))
top_marks = sorted(MARKS, key=lambda s: s[2], reverse=True)
grade_7 = [s for s in MARKS if s[1] == PICK]


def build_tuition_marks(wb, finished):
    ws = wb.active
    put_rows(ws, [HEAD2, *MARKS])
    put_rows(ws, [('Grades',)], start_col=5)
    put_rows(ws, [HEAD2], start_col=7)
    put_rows(ws, [HEAD2], start_col=11)
    fit_widths(ws)
    if finished:
        array_formula(ws, 'E2', len(grades), 1, '=SORT(UNIQUE(B2:B9))')
        array_formula(ws, 'G2', len(top_marks), 3, '=SORT(A2:C9,3,-1)')
        array_formula(ws, 'K2', len(grade_7), 3, '=FILTER(A2:C9,B2:B9="Grade 7","No students")')


marks_expected = {}
put_expected(marks_expected, 'E', 2, [(c,) for c in grades])
put_expected(marks_expected, 'G', 2, top_marks)
put_expected(marks_expected, 'K', 2, grade_7)

FILES = [
    XlsxFile('shop-sales', build_shop_sales, sales_expected, fit='none'),
    XlsxFile('tuition-marks', build_tuition_marks, marks_expected, fit='none'),
]

# The lesson states these results; keep them in step with the page.
assert products == ['Sugar', 'Tea', 'Rice', 'Cooking oil']
assert [s[2] for s in by_amount] == [3100, 2700, 2400, 1800, 1500, 1300, 1200, 950, 900, 850]
assert by_amount[0] == ('Wednesday', 'Cooking oil', 3100)
assert by_amount[-1] == ('Monday', 'Tea', 850)
assert big_sales == [('Tuesday', 'Rice', 2400), ('Wednesday', 'Cooking oil', 3100),
                     ('Thursday', 'Cooking oil', 2700)]
assert len({s[2] for s in SALES}) == len(SALES)   # no ties, so the sorted order is certain
assert unique([c for _, c, _ in MARKS]) == ['Grade 8', 'Grade 7', 'Grade 6']
assert grades == ['Grade 6', 'Grade 7', 'Grade 8']
assert [s[0] for s in top_marks] == ['Amna', 'Noor', 'Bilal', 'Sana', 'Hassan', 'Usman',
                                     'Ali Raza', 'Iqra']
assert len({s[2] for s in MARKS}) == len(MARKS)
assert grade_7 == [('Noor', 'Grade 7', 91), ('Usman', 'Grade 7', 72), ('Amna', 'Grade 7', 95)]

# Step by Step step 9 changes C10 to 2200 for a moment (then back to 900).
_changed = [(d, p, 2200 if (d, p) == ('Friday', 'Sugar') else a) for d, p, a in SALES]
assert [s for s in _changed if s[2] > OVER][-1] == ('Friday', 'Sugar', 2200)
assert len([s for s in _changed if s[2] > OVER]) == 4
assert sorted(_changed, key=lambda s: s[2], reverse=True)[3] == ('Friday', 'Sugar', 2200)
