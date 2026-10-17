"""Lesson 7.27: Conditional Formatting.

Conditional formatting has no formulas in the cells, so there are no answer cells to recalculate
(`expected` is empty). The finished files hold the same rules the lesson makes; the asserts below
check the rules and every cell list the lesson states.

The colours are Excel's usual highlight colours: light red fill with dark red text, green fill
with dark green text, and yellow fill with dark yellow text.
"""

from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle

from xlsxkit import XlsxFile, new_workbook, put_rows

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'conditional-formatting.mdx'
FOLDER = 'conditional-formatting'

RED = dict(fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
           font=Font(color='9C0006'))
GREEN = dict(fill=PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'),
             font=Font(color='006100'))
YELLOW = dict(fill=PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid'),
              font=Font(color='9C5700'))

# Step by Step: a tuition centre's test marks and fee status.
MARKS = [
    ('Ali', 78, 35, 62, 'Paid'), ('Sana', 92, 88, 95, 'Paid'), ('Hamza', 38, 54, 41, 'Unpaid'),
    ('Fatima', 65, 71, 90, 'Paid'), ('Usman', 45, 29, 33, 'Unpaid'), ('Ayesha', 91, 77, 84, 'Paid'),
    ('Bilal', 52, 61, 39, 'Paid'), ('Hira', 70, 83, 58, 'Unpaid'),
]
HEADINGS = ('Student', 'Maths', 'English', 'Science', 'Fee')
MARK_RANGE = 'B2:D9'
FEE_RANGE = 'E2:E9'
LOW, HIGH = 40, 89          # Less Than 40; Greater Than 89 (90 and above)
CHANGED_CELL, CHANGED_TO = 'C6', 45   # step 13: Usman's English mark is corrected


def marks_after_change():
    rows = [list(r) for r in MARKS]
    rows[4][2] = CHANGED_TO            # row 6 = Usman, column C = English
    return [tuple(r) for r in rows]


def build_test_marks(wb, finished):
    ws = wb.active
    put_rows(ws, [HEADINGS, *(marks_after_change() if finished else MARKS)])
    if finished:
        ws.conditional_formatting.add(MARK_RANGE, CellIsRule(operator='lessThan',
                                                             formula=[str(LOW)], **RED))
        ws.conditional_formatting.add(MARK_RANGE, CellIsRule(operator='greaterThan',
                                                             formula=[str(HIGH)], **GREEN))
        ws.conditional_formatting.add(FEE_RANGE, Rule(
            type='containsText', operator='containsText', text='Unpaid',
            formula=['NOT(ISERROR(SEARCH("Unpaid",E2)))'], dxf=DifferentialStyle(**YELLOW)))


# Practice: a grocery shop's stock list. Low stock = fewer than 10 in stock. The task also tries
# Data Bars on column C and then clears them, so the finished file has only the low-stock rule.
STOCK = [('Sugar (kg)', 25, 60), ('Tea (pack)', 8, 45), ('Cooking oil (litre)', 4, 38),
         ('Rice (kg)', 30, 52), ('Daal (kg)', 12, 20), ('Atta (10 kg bag)', 6, 30),
         ('Salt (pack)', 18, 15), ('Soap', 9, 25), ('Biscuits (pack)', 40, 70),
         ('Matchbox', 50, 12)]
STOCK_RANGE = 'B2:B11'
LOW_STOCK = 10


def build_shop_stock(wb, finished):
    ws = wb.active
    put_rows(ws, [('Item', 'In Stock', 'Sold This Month'), *STOCK])
    if finished:
        ws.conditional_formatting.add(STOCK_RANGE, CellIsRule(operator='lessThan',
                                                              formula=[str(LOW_STOCK)], **RED))


FILES = [
    XlsxFile('test-marks', build_test_marks),
    XlsxFile('shop-stock', build_shop_stock),
]


# The lesson states these cells; keep them in step with the page.
def cells(rows, test):
    found = []
    for r, row in enumerate(rows, start=2):
        for c, letter in zip(row[1:4], 'BCD'):
            if test(c):
                found.append(f'{letter}{r}')
    return found


assert cells(MARKS, lambda m: m < LOW) == ['C2', 'B4', 'C6', 'D6', 'D8']           # step 6
assert cells(MARKS, lambda m: m > HIGH) == ['B3', 'D3', 'D5', 'B7']                 # step 9
assert [f'E{r}' for r, row in enumerate(MARKS, start=2) if 'Unpaid' in row[4]] == [
    'E4', 'E6', 'E9']                                                               # step 12
assert MARKS[4][0] == 'Usman' and MARKS[4][2] == 29                                 # step 13
assert cells(marks_after_change(), lambda m: m < LOW) == ['C2', 'B4', 'D6', 'D8']  # step 14
assert [(r[0], r[1]) for r in STOCK if r[1] < LOW_STOCK] == [
    ('Tea (pack)', 8), ('Cooking oil (litre)', 4), ('Atta (10 kg bag)', 6), ('Soap', 9)]
assert [f'B{r}' for r, row in enumerate(STOCK, start=2) if row[1] < LOW_STOCK] == [
    'B3', 'B4', 'B7', 'B9']                                                         # Practice 3
# The Between backup (0 to 39, 0 to 9) colours the same cells, because all values are whole.
assert all(isinstance(m, int) for r in MARKS for m in r[1:4])
assert cells(MARKS, lambda m: 0 <= m <= 39) == cells(MARKS, lambda m: m < LOW)
assert [r for r in STOCK if 0 <= r[1] <= 9] == [r for r in STOCK if r[1] < LOW_STOCK]
assert [m for r in MARKS for m in r[1:4] if m in (40, 89, 90)] == [90]          # edge values
assert max(STOCK, key=lambda r: r[2])[:3:2] == ('Biscuits (pack)', 70)               # longest bar
assert min(STOCK, key=lambda r: r[2])[:3:2] == ('Matchbox', 12)                      # shortest bar


def _rules(ws):
    return [(str(cf.sqref), rule.type, rule.operator, list(rule.formula))
            for cf in ws.conditional_formatting for rule in cf.rules]


_wb = new_workbook()
build_test_marks(_wb, True)
assert sorted(_rules(_wb.active)) == sorted([
    (MARK_RANGE, 'cellIs', 'lessThan', ['40']),
    (MARK_RANGE, 'cellIs', 'greaterThan', ['89']),
    (FEE_RANGE, 'containsText', 'containsText', ['NOT(ISERROR(SEARCH("Unpaid",E2)))']),
])
assert _wb.active[CHANGED_CELL].value == CHANGED_TO
_wb = new_workbook()
build_shop_stock(_wb, True)
assert _rules(_wb.active) == [(STOCK_RANGE, 'cellIs', 'lessThan', ['10'])]
for _build in (build_test_marks, build_shop_stock):
    _wb = new_workbook()
    _build(_wb, False)
    assert _rules(_wb.active) == []
