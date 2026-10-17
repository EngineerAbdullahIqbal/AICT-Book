"""Lesson 7.13: Project: Student Result Sheet.

One table, used in Step by Step and Practice. Starter: names and marks in the typing order,
columns F to H empty. Finished: the end of the Practice, as in the lesson's answer key: rows
sorted by Total (Largest to Smallest), Total/Average/Percentage formulas, Highest and Lowest
rows, the "Total Marks" column chart, and the page set to Landscape, 1 page by 1 page.
"""

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef

from xlsxkit import XlsxFile, put_rows, formula

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-result-sheet.mdx'
FOLDER = 'project-result-sheet'

HEADINGS = ('Name', 'English', 'Urdu', 'Maths', 'Computer', 'Total', 'Average', 'Percentage')
STUDENTS = [('Ayesha Khan', 42, 45, 38, 47), ('Bilal Ahmed', 35, 40, 44, 39),
            ('Fatima Noor', 48, 46, 41, 45), ('Hamza Ali', 30, 36, 33, 28),
            ('Sana Riaz', 44, 39, 47, 43), ('Usman Tariq', 38, 42, 29, 35)]

# Practice step 3: sort by Total, Largest to Smallest (each row moves together).
SORTED = sorted(STUDENTS, key=lambda s: sum(s[1:]), reverse=True)


def build_result_sheet(wb, finished):
    ws = wb.active
    put_rows(ws, [HEADINGS, *(SORTED if finished else STUDENTS)])
    if not finished:
        return
    for r in range(2, 8):
        formula(ws, f'F{r}', f'=SUM(B{r}:E{r})')
        formula(ws, f'G{r}', f'=AVERAGE(B{r}:E{r})')
        formula(ws, f'H{r}', f'=F{r}/200*100')
    ws['A9'] = 'Highest'
    ws['A10'] = 'Lowest'
    for col in 'BCDEF':
        formula(ws, f'{col}9', f'=MAX({col}2:{col}7)')
        formula(ws, f'{col}10', f'=MIN({col}2:{col}7)')
    chart = BarChart()
    chart.type = 'col'
    chart.grouping = 'clustered'
    chart.title = 'Total Marks'
    chart.add_data(Reference(ws, min_col=6, min_row=1, max_row=7), titles_from_data=True)
    # The names are text, so the categories are a text reference (strRef).
    chart.series[0].cat = AxDataSource(strRef=StrRef(f=f"'{ws.title}'!$A$2:$A$7"))
    chart.legend = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    ws.add_chart(chart, 'J2')
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


expected = {}
for r, s in zip(range(2, 8), SORTED):
    total = sum(s[1:])
    expected[f'F{r}'] = total
    expected[f'G{r}'] = total / 4
    expected[f'H{r}'] = total / 200 * 100
for i, col in enumerate('BCDEF'):
    column = [sum(s[1:]) if col == 'F' else s[1 + i] for s in SORTED]
    expected[f'{col}9'] = max(column)
    expected[f'{col}10'] = min(column)

FILES = [XlsxFile('result-sheet', build_result_sheet, expected)]

# The lesson's answer key; keep it in step with the page.
ANSWER_KEY = [
    ('Fatima Noor', 48, 46, 41, 45, 180, 45, 90),
    ('Sana Riaz', 44, 39, 47, 43, 173, 43.25, 86.5),
    ('Ayesha Khan', 42, 45, 38, 47, 172, 43, 86),
    ('Bilal Ahmed', 35, 40, 44, 39, 158, 39.5, 79),
    ('Usman Tariq', 38, 42, 29, 35, 144, 36, 72),
    ('Hamza Ali', 30, 36, 33, 28, 127, 31.75, 63.5),
]
for r, (s, key) in enumerate(zip(SORTED, ANSWER_KEY), start=2):
    assert s == key[:5]
    assert expected[f'F{r}'] == key[5] and expected[f'G{r}'] == key[6]
    assert abs(expected[f'H{r}'] - key[7]) < 1e-9
assert [expected[f'{c}9'] for c in 'BCDEF'] == [48, 46, 47, 47, 180]
assert [expected[f'{c}10'] for c in 'BCDEF'] == [30, 36, 29, 28, 127]
# Step by Step results for Ayesha (before sorting): 172, 43, 86.
assert sum(STUDENTS[0][1:]) == 172 and sum(STUDENTS[0][1:]) / 4 == 43
assert sum(STUDENTS[0][1:]) / 200 * 100 == 86
