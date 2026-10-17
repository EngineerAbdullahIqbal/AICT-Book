"""Lesson 7.29: Project: Shop Stock and Sales Register.

Two registers with the same layout (A Item, B Price (Rs), C Opening Stock, D Sold,
E Closing Stock, F Sale Value (Rs), G Status), items in rows 2-9, row 10 empty, totals in row 11.
Row 10 stays empty so that the filter covers only the items (A1:G9), not the totals.

Starter: data only (A-D) and the Total label. Finished: the end of the task: the formulas,
the "Order more" highlight rule on G2:G9, and the filter on Status showing only "Order more"
(openpyxl stores the filter but does not hide rows, so build() hides them, as Excel does).
"""

from openpyxl.formatting.rule import Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle

from xlsxkit import XlsxFile, new_workbook, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-stock-and-sales.mdx'
FOLDER = 'project-stock-and-sales'

HEADINGS = ('Item', 'Price (Rs)', 'Opening Stock', 'Sold', 'Closing Stock', 'Sale Value (Rs)',
            'Status')
ALERT = 'Order more'
RED = DifferentialStyle(fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE',
                                         fill_type='solid'), font=Font(color='9C0006'))

# Step by Step: Kashif's mobile accessories shop, October. Order more when fewer than 5 are left.
MOBILE = [
    ('Phone charger', 450, 30, 26), ('Data cable', 250, 40, 22), ('Earphones', 350, 25, 21),
    ('Screen protector', 150, 60, 48), ('Phone cover', 300, 35, 19), ('Power bank', 2200, 10, 7),
    ('Memory card 32 GB', 800, 15, 9), ('Car charger', 500, 12, 7),
]
MOBILE_LIMIT = 5

# Practice: Shabana's general store, October. Order more when fewer than 10 are left.
# Practice step: two more packs of biscuits are sold, so D8 changes from 70 to 72.
STORE = [
    ('Sugar (kg)', 155, 60, 52), ('Cooking oil (litre)', 545, 40, 26), ('Tea (pack)', 480, 30, 24),
    ('Rice (kg)', 340, 50, 35), ('Daal (kg)', 285, 25, 18), ('Soap (bar)', 95, 48, 30),
    ('Biscuits (pack)', 60, 80, 70), ('Matchbox', 10, 100, 64),
]
STORE_LIMIT = 10
STORE_CHANGED = [r if r[0] != 'Biscuits (pack)' else (r[0], r[1], r[2], 72) for r in STORE]

LAST = 9          # last item row
TOTAL_ROW = 11


def make_builder(rows_start, rows_end, limit):
    def build(wb, finished):
        ws = wb.active
        rows = rows_end if finished else rows_start
        put_rows(ws, [HEADINGS, *rows])
        ws[f'A{TOTAL_ROW}'] = 'Total'
        bold(ws, f'A{TOTAL_ROW}')
        if not finished:
            return
        for r in range(2, LAST + 1):
            formula(ws, f'E{r}', f'=C{r}-D{r}')
            formula(ws, f'F{r}', f'=D{r}*B{r}')
            formula(ws, f'G{r}', f'=IF(E{r}<{limit},"{ALERT}","OK")')
        for col in 'CDEF':
            formula(ws, f'{col}{TOTAL_ROW}', f'=SUM({col}2:{col}{LAST})')
        ws.conditional_formatting.add(f'G2:G{LAST}', Rule(
            type='containsText', operator='containsText', text=ALERT,
            formula=[f'NOT(ISERROR(SEARCH("{ALERT}",G2)))'], dxf=RED))
        ws.auto_filter.ref = f'A1:G{LAST}'
        ws.auto_filter.add_filter_column(6, [ALERT])            # column G (Status)
        for r, row in enumerate(rows, start=2):
            if not row[2] - row[3] < limit:
                ws.row_dimensions[r].hidden = True
    return build


def expected_for(rows, limit):
    exp = {}
    for r, (_, price, opening, sold) in enumerate(rows, start=2):
        exp[f'E{r}'] = opening - sold
        exp[f'F{r}'] = sold * price
        exp[f'G{r}'] = ALERT if opening - sold < limit else 'OK'
    exp[f'C{TOTAL_ROW}'] = sum(r[2] for r in rows)
    exp[f'D{TOTAL_ROW}'] = sum(r[3] for r in rows)
    exp[f'E{TOTAL_ROW}'] = sum(r[2] - r[3] for r in rows)
    exp[f'F{TOTAL_ROW}'] = sum(r[1] * r[3] for r in rows)
    return exp


mobile = expected_for(MOBILE, MOBILE_LIMIT)
store = expected_for(STORE_CHANGED, STORE_LIMIT)
store_before = expected_for(STORE, STORE_LIMIT)

FILES = [
    XlsxFile('mobile-shop-stock', make_builder(MOBILE, MOBILE, MOBILE_LIMIT), mobile),
    XlsxFile('general-store-stock', make_builder(STORE, STORE_CHANGED, STORE_LIMIT), store),
]

# The lesson states these results; keep them in step with the page.
assert [mobile[f'E{r}'] for r in range(2, 10)] == [4, 18, 4, 12, 16, 3, 6, 5]
assert [mobile[f'F{r}'] for r in range(2, 10)] == [11700, 5500, 7350, 7200, 5700, 15400, 7200,
                                                    3500]
assert [mobile[f'G{r}'] for r in range(2, 10)] == [
    ALERT, 'OK', ALERT, 'OK', 'OK', ALERT, 'OK', 'OK']
assert (mobile['C11'], mobile['D11'], mobile['E11'], mobile['F11']) == (227, 159, 68, 63550)
assert mobile['E2'] == 4 and mobile['F2'] == 11700                       # step 4 and step 6

assert [store_before[f'E{r}'] for r in range(2, 10)] == [8, 14, 6, 15, 7, 18, 10, 36]
assert [store_before[f'G{r}'] for r in range(2, 10)].count(ALERT) == 3
assert store_before['F11'] == 58470
assert [store[f'E{r}'] for r in range(2, 10)] == [8, 14, 6, 15, 7, 18, 8, 36]
assert [store[f'F{r}'] for r in range(2, 10)] == [8060, 14170, 11520, 11900, 5130, 2850, 4320,
                                                   640]
assert [store[f'G{r}'] for r in range(2, 10)] == [
    ALERT, 'OK', ALERT, 'OK', ALERT, 'OK', ALERT, 'OK']
assert (store['C11'], store['D11'], store['E11'], store['F11']) == (433, 321, 112, 58590)

# The finished files: the filter shows only the "Order more" rows; the starters have no filter.
for _rows, _end, _limit, _shown in ((MOBILE, MOBILE, MOBILE_LIMIT, [2, 4, 7]),
                                    (STORE, STORE_CHANGED, STORE_LIMIT, [2, 4, 6, 8])):
    _wb = new_workbook()
    make_builder(_rows, _end, _limit)(_wb, True)
    _ws = _wb.active
    assert _ws.auto_filter.ref == 'A1:G9'
    assert [r for r in range(2, 10) if not _ws.row_dimensions[r].hidden] == _shown
    assert len(_ws.conditional_formatting) == 1
    _wb = new_workbook()
    make_builder(_rows, _end, _limit)(_wb, False)
    assert _wb.active.auto_filter.ref is None and len(_wb.active.conditional_formatting) == 0
