"""Lesson 7.32: Project: Staff Attendance Sheet.

Two attendance sheets with the same layout: A Name, then one column per working day (the
heading is the day of the month, as a number), then Present, Absent, Leave and Attendance %.
Each day cell holds a code: P (present), A (absent) or L (leave). Staff start in row 2.
Present/Absent/Leave = COUNTIF(day cells, "P"/"A"/"L");
Attendance % = ROUND(Present / COUNTA(day cells) * 100, 0).

Starter: names and codes only. Finished: the formulas, a "Text that Contains A" rule (red) on
the day cells, and the panes frozen at B2 (row 1 and column A stay in view).
The Practice sheet also has an Absent row (absent people per day, COUNTIF down each column).
Its finished file shows the end of the task, after Kiran's absence on 1 December is changed to L.
"""

from decimal import Decimal, ROUND_HALF_UP

from openpyxl.formatting.rule import Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter

from xlsxkit import XlsxFile, new_workbook, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-attendance-sheet.mdx'
FOLDER = 'project-attendance-sheet'

RED = DifferentialStyle(fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE',
                                         fill_type='solid'), font=Font(color='9C0006'))

# Step by Step: a courier office in Gujranwala, working days 2-18 November 2026 (Monday to
# Saturday; 8 and 15 November are Sundays).
OFFICE_DAYS = [2, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13, 14, 16, 17, 18]
OFFICE = [
    ('Imran Haider', 'PPPPPPPPPPPPPPP'),
    ('Saba Naz', 'PPPAPPPPPPLPPPP'),
    ('Kashif Mehmood', 'PPAPPPPAPPPPAPP'),
    ('Nadia Parveen', 'PPPPPPLPPPPPPPP'),
    ('Waqar Ahmed', 'APPPAPLLPPPAPPP'),
    ('Shazia Bibi', 'PPPPLLLPPPPPAPP'),
    ('Adnan Rafiq', 'PPPPPPPPPPPPPPP'),
    ('Mehwish Tariq', 'PPPPPPPPPAPPPPP'),
    ('Zahid Hussain', 'PAPAPPLAPPAPPAP'),
    ('Rabia Anwar', 'PPPPPPPPLPPPAPP'),
]

# Practice: a hardware shop in Sialkot, working days 1-14 December 2026 (6 and 13 are Sundays).
SHOP_DAYS = [1, 2, 3, 4, 5, 7, 8, 9, 10, 11, 12, 14]
SHOP = [
    ('Arshad Ali', 'PPPPPPPPPPPP'),
    ('Sajida Bano', 'PPAPPPPPPLPP'),
    ('Nasir Khan', 'PPPPAPPPAPPA'),
    ('Farah Naz', 'PLLPPPPPPPPP'),
    ('Bilal Hussain', 'PPPPPAPPPPPP'),
    ('Kiran Shahzad', 'APPPPPPPAPPP'),
]
# Practice step: Kiran's absence on 1 December was approved leave, so B7 changes from A to L.
SHOP_CHANGED = [(n, ('L' + c[1:]) if n == 'Kiran Shahzad' else c) for n, c in SHOP]


def round_half_up(x, digits=0):
    return float(Decimal(repr(x)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))


def layout(days):
    first_day = 'B'
    last_day = get_column_letter(1 + len(days))
    cols = [get_column_letter(2 + len(days) + k) for k in range(4)]   # Present, Absent, Leave, %
    return first_day, last_day, cols


def make_builder(days, start_rows, end_rows, absent_row=None):
    first_day, last_day, (cp, ca, cl, cpct) = layout(days)
    headings = ('Name', *days, 'Present', 'Absent', 'Leave', 'Attendance %')

    def build(wb, finished):
        ws = wb.active
        rows = end_rows if finished else start_rows
        put_rows(ws, [headings, *((name, *codes) for name, codes in rows)])
        last = 1 + len(rows)
        if absent_row:
            ws[f'A{absent_row}'] = 'Absent'
            bold(ws, f'A{absent_row}')
        if not finished:
            return
        for r in range(2, last + 1):
            days_ref = f'{first_day}{r}:{last_day}{r}'
            formula(ws, f'{cp}{r}', f'=COUNTIF({days_ref},"P")')
            formula(ws, f'{ca}{r}', f'=COUNTIF({days_ref},"A")')
            formula(ws, f'{cl}{r}', f'=COUNTIF({days_ref},"L")')
            formula(ws, f'{cpct}{r}', f'=ROUND({cp}{r}/COUNTA({days_ref})*100,0)')
        if absent_row:
            for k in range(len(days)):
                col = get_column_letter(2 + k)
                formula(ws, f'{col}{absent_row}', f'=COUNTIF({col}2:{col}{last},"A")')
        ws.conditional_formatting.add(f'{first_day}2:{last_day}{last}', Rule(
            type='containsText', operator='containsText', text='A',
            formula=[f'NOT(ISERROR(SEARCH("A",{first_day}2)))'], dxf=RED))
        ws.freeze_panes = 'B2'
    return build


def expected_for(days, rows, absent_row=None):
    _, _, (cp, ca, cl, cpct) = layout(days)
    exp = {}
    for r, (_, codes) in enumerate(rows, start=2):
        assert len(codes) == len(days) and set(codes) <= set('PAL')
        p, a, l = codes.count('P'), codes.count('A'), codes.count('L')
        exp[f'{cp}{r}'], exp[f'{ca}{r}'], exp[f'{cl}{r}'] = p, a, l
        exp[f'{cpct}{r}'] = round_half_up(p / len(codes) * 100)
    if absent_row:
        for k in range(len(days)):
            exp[f'{get_column_letter(2 + k)}{absent_row}'] = sum(c[k] == 'A' for _, c in rows)
    return exp


SHOP_ABSENT_ROW = 9
office = expected_for(OFFICE_DAYS, OFFICE)
shop_before = expected_for(SHOP_DAYS, SHOP, SHOP_ABSENT_ROW)
shop = expected_for(SHOP_DAYS, SHOP_CHANGED, SHOP_ABSENT_ROW)

FILES = [
    XlsxFile('office-attendance', make_builder(OFFICE_DAYS, OFFICE, OFFICE), office),
    XlsxFile('shop-attendance', make_builder(SHOP_DAYS, SHOP, SHOP_CHANGED, SHOP_ABSENT_ROW),
             shop),
]


# The lesson states these results; keep them in step with the page.
def col(exp, c, first, last):
    return [exp[f'{c}{r}'] for r in range(first, last + 1)]


assert layout(OFFICE_DAYS) == ('B', 'P', ['Q', 'R', 'S', 'T'])
assert col(office, 'Q', 2, 11) == [15, 13, 12, 14, 10, 11, 15, 14, 9, 13]
assert col(office, 'R', 2, 11) == [0, 1, 3, 0, 3, 1, 0, 1, 5, 1]
assert col(office, 'S', 2, 11) == [0, 1, 0, 1, 2, 3, 0, 0, 1, 1]
assert col(office, 'T', 2, 11) == [100, 87, 80, 93, 67, 73, 100, 93, 60, 87]
assert sum(col(office, 'R', 2, 11)) == 15

assert layout(SHOP_DAYS) == ('B', 'M', ['N', 'O', 'P', 'Q'])
assert col(shop_before, 'O', 2, 7) == [0, 1, 3, 0, 1, 2]
assert shop_before['B9'] == 1
assert col(shop, 'N', 2, 7) == [12, 10, 9, 10, 11, 10]
assert col(shop, 'O', 2, 7) == [0, 1, 3, 0, 1, 1]
assert col(shop, 'P', 2, 7) == [0, 1, 0, 2, 0, 1]
assert col(shop, 'Q', 2, 7) == [100, 83, 75, 83, 92, 83]
assert [shop[f'{get_column_letter(2 + k)}9'] for k in range(12)] == [
    0, 0, 1, 0, 1, 1, 0, 0, 2, 0, 0, 1]

# The finished files are frozen at B2 and have one colour rule; the starters have neither.
for _days, _start, _end, _absent in ((OFFICE_DAYS, OFFICE, OFFICE, None),
                                     (SHOP_DAYS, SHOP, SHOP_CHANGED, SHOP_ABSENT_ROW)):
    for _finished in (False, True):
        _wb = new_workbook()
        make_builder(_days, _start, _end, _absent)(_wb, _finished)
        _ws = _wb.active
        assert _ws.freeze_panes == ('B2' if _finished else None)
        assert len(_ws.conditional_formatting) == (1 if _finished else 0)

# Red cells (one per A code): 15 in the office sheet, 6 in the shop sheet after the change.
assert sum(c.count('A') for _, c in OFFICE) == 15 and sum(c.count('A') for _, c in SHOP_CHANGED) == 6
# Zahid's absences are in C10, E10, I10, L10 and O10.
assert [get_column_letter(2 + k) for k, c in enumerate(OFFICE[8][1]) if c == 'A'] == list('CEILO')
