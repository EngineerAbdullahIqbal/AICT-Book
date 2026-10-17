"""Lesson 7.30: Project: Customer Bill With Discount.

Two bills with the same layout:
  A1 shop name; A2:B5 Bill No., Date, Customer, Discount (rate in B5);
  row 7 headings (A Item, B Qty, C Rate (Rs), D Amount (Rs), E Discount (Rs)); items in rows 8-13;
  C15:D17 Subtotal, Discount and Grand Total.
Amount = Qty x Rate; Discount = ROUND(Amount x $B$5, 0); Subtotal = SUM of amounts;
Discount total = SUM of discounts; Grand Total = Subtotal - Discount.

Starter: the bill details and items (data only); the rate is typed as a plain number (0.05), so
the student applies Percent Style. Finished: the formulas and the number formats the lesson
applies (Percent Style, Short Date, Number with 0 decimal places and the 1000 separator), bold
headings and borders. The Practice finished file shows the end of the task, with the rate
changed to 12%.
"""

import math
from datetime import datetime

from openpyxl.styles import Border, Font, Side

from xlsxkit import XlsxFile, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-customer-bill.mdx'
FOLDER = 'project-customer-bill'

HEADINGS = ('Item', 'Qty', 'Rate (Rs)', 'Amount (Rs)', 'Discount (Rs)')
FIRST, LAST = 8, 13
SHORT_DATE = 'mm-dd-yy'        # Excel's built-in format 14 = Short Date (follows the PC's setting)
MONEY = '#,##0'                # Number, 0 decimal places, Use 1000 Separator
THIN = Side(style='thin')

# Step by Step: Rafiq General Store, a sale week with 5% off every item.
GROCERY = dict(
    shop='Rafiq General Store', bill_no=2045, date=datetime(2026, 10, 15),
    customer='Nasreen Akhtar', rate=0.05,
    items=[('Basmati rice (kg)', 5, 340), ('Sugar (kg)', 3, 155), ('Cooking oil (litre)', 3, 545),
           ('Tea (pack)', 1, 480), ('Daal (kg)', 4, 285), ('Washing powder (pack)', 1, 365)],
)
# Practice: Bright Stationers, a back-to-school offer of 10%, then 12% for a regular customer.
STATIONERY = dict(
    shop='Bright Stationers', bill_no=318, date=datetime(2026, 10, 16),
    customer='Kamran Yousaf', rate=0.10,
    items=[('Notebook', 12, 85), ('Ball pens (pack of 10)', 2, 245), ('Geometry box', 1, 420),
           ('School bag', 1, 2450), ('Eraser', 6, 18), ('Sharpener', 6, 23)],
)
STATIONERY_NEW_RATE = 0.12


def make_builder(bill, final_rate):
    def build(wb, finished):
        ws = wb.active
        ws['A1'] = bill['shop']
        ws['A1'].font = Font(bold=True, size=16)
        put_rows(ws, [('Bill No.', bill['bill_no']), ('Date', bill['date']),
                      ('Customer', bill['customer']),
                      ('Discount', final_rate if finished else bill['rate'])],
                 start_row=2, bold_first=False)
        bold(ws, 'A2', 'A3', 'A4', 'A5')
        ws['B3'].number_format = SHORT_DATE
        put_rows(ws, [HEADINGS, *bill['items']], start_row=7)
        put_rows(ws, [('Subtotal',), ('Discount',), ('Grand Total',)], start_row=15, start_col=3,
                 bold_first=False)
        bold(ws, 'C15', 'C16', 'C17')
        if not finished:
            return
        for r in range(FIRST, LAST + 1):
            formula(ws, f'D{r}', f'=B{r}*C{r}')
            formula(ws, f'E{r}', f'=ROUND(D{r}*$B$5,0)')
        formula(ws, 'D15', f'=SUM(D{FIRST}:D{LAST})')
        formula(ws, 'D16', f'=SUM(E{FIRST}:E{LAST})')
        formula(ws, 'D17', '=D15-D16')
        bold(ws, 'D17')
        ws['B5'].number_format = '0%'
        for r in range(FIRST, LAST + 1):
            for col in 'CDE':
                ws[f'{col}{r}'].number_format = MONEY
        for ref in ('D15', 'D16', 'D17'):
            ws[ref].number_format = MONEY
        for row in ws[f'A7:E{LAST}']:
            for cell in row:
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    return build


def excel_round0(x):
    """ROUND(x, 0) for the positive numbers here: halves round up, as in Excel."""
    assert abs(x - math.floor(x) - 0.5) > 1e-6, 'avoid exact halves in practice data'
    return math.floor(x + 0.5)


def expected_for(bill, rate):
    exp = {}
    for r, (_, qty, price) in enumerate(bill['items'], start=FIRST):
        exp[f'D{r}'] = qty * price
        exp[f'E{r}'] = excel_round0(qty * price * rate)
    exp['D15'] = sum(exp[f'D{r}'] for r in range(FIRST, LAST + 1))
    exp['D16'] = sum(exp[f'E{r}'] for r in range(FIRST, LAST + 1))
    exp['D17'] = exp['D15'] - exp['D16']
    return exp


grocery = expected_for(GROCERY, GROCERY['rate'])
stationery_10 = expected_for(STATIONERY, STATIONERY['rate'])
stationery = expected_for(STATIONERY, STATIONERY_NEW_RATE)

FILES = [
    XlsxFile('grocery-bill', make_builder(GROCERY, GROCERY['rate']), grocery),
    XlsxFile('stationery-bill', make_builder(STATIONERY, STATIONERY_NEW_RATE), stationery),
]

# The lesson states these results; keep them in step with the page.
_rows = range(FIRST, LAST + 1)
assert [grocery[f'D{r}'] for r in _rows] == [1700, 465, 1635, 480, 1140, 365]
# Before ROUND (step 6): the discounts with decimals.
assert [round(a * 0.05, 2) for a in [1700, 465, 1635, 480, 1140, 365]] == [
    85, 23.25, 81.75, 24, 57, 18.25]
assert [grocery[f'E{r}'] for r in _rows] == [85, 23, 82, 24, 57, 18]
assert (grocery['D15'], grocery['D16'], grocery['D17']) == (5785, 289, 5496)

assert [stationery_10[f'D{r}'] for r in _rows] == [1020, 490, 420, 2450, 108, 138]
assert [stationery_10[f'E{r}'] for r in _rows] == [102, 49, 42, 245, 11, 14]
assert (stationery_10['D15'], stationery_10['D16'], stationery_10['D17']) == (4626, 463, 4163)
assert [stationery[f'E{r}'] for r in _rows] == [122, 59, 50, 294, 13, 17]
assert (stationery['D15'], stationery['D16'], stationery['D17']) == (4626, 555, 4071)
