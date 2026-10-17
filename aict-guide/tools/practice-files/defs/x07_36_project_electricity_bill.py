"""Lesson 7.36: Project: Electricity Bill Calculator.

THE RATES ARE MADE UP FOR PRACTICE. They are not the real electricity tariff (owner decision,
specs/practical-upgrade/spec.md §4.2 and PU-05). Row 1 of every file says so.

Two sheets with the same layout:
  A1 the practice-rates note (merged across A1:H1);
  A3:B6 the rates, each in one cell: B3 low rate (Rs per unit), B4 high rate (Rs per unit),
  B5 low-rate limit (units), B6 fixed charge (Rs);
  row 8 headings: A Flat / Month, B Old Reading, C New Reading, D Units Used, E Low-Rate Units,
  F High-Rate Units, G Energy Cost (Rs), H Total Bill (Rs); data from row 9; a Total row below.
Units Used = C-B; Low-Rate Units = IF(D>$B$5,$B$5,D); High-Rate Units = D-E;
Energy Cost = E*$B$3+F*$B$4; Total Bill = ROUND(G+$B$6,0). The Total row adds D to H.
Practice only: A15 July minus October (Rs), B15 =H9-H12.

Starter: the note, rates and meter readings (data only). Finished: the formulas. The Practice
finished file shows the end of the task, with the high rate changed to 30.6.
"""

import math

from openpyxl.styles import Font

from xlsxkit import XlsxFile, put_rows, formula, bold

APP = 'excel'
MODULE = 'module-07-ms-excel'
LESSON_FILE = 'project-electricity-bill.mdx'
FOLDER = 'project-electricity-bill'

NOTE = ('PRACTICE RATES ONLY: these rates are made up for learning. '
        'They are NOT the real electricity tariff.')
LABELS = ('Low rate (Rs per unit)', 'High rate (Rs per unit)', 'Low-rate limit (units)',
          'Fixed charge (Rs)')
HEADINGS = ('Old Reading', 'New Reading', 'Units Used', 'Low-Rate Units', 'High-Rate Units',
            'Energy Cost (Rs)', 'Total Bill (Rs)')
FIRST = 9
RATES = (18.4, 27.9, 100, 150)       # made-up practice rates

# Step by Step: five flats in one building in Rawalpindi, each with its own sub-meter.
FLATS = dict(first_heading='Flat', rates=RATES, rows=[
    ('Flat 1', 4520, 4607), ('Flat 2', 3815, 3963), ('Flat 3', 6102, 6202),
    ('Flat 4', 2740, 2953), ('Flat 5', 5388, 5449)])
# Practice: one family's house, July to October. Each old reading is last month's new reading.
# Later the practice high rate goes up from 27.9 to 30.6 (one cell, B4).
MONTHS = dict(first_heading='Month', rates=RATES, compare=True, rows=[
    ('July', 8240, 8506), ('August', 8506, 8738), ('September', 8738, 8917),
    ('October', 8917, 9011)])
NEW_HIGH_RATE = 30.6


def make_builder(sheet, final_rates):
    rows = sheet['rows']
    last = FIRST + len(rows) - 1
    total = last + 1

    def build(wb, finished):
        ws = wb.active
        ws['A1'] = NOTE
        ws['A1'].font = Font(bold=True, color='C00000')
        ws.merge_cells('A1:H1')
        rates = final_rates if finished else sheet['rates']
        put_rows(ws, list(zip(LABELS, rates)), start_row=3, bold_first=False)
        bold(ws, 'A3', 'A4', 'A5', 'A6')
        put_rows(ws, [(sheet['first_heading'], *HEADINGS), *rows], start_row=FIRST - 1)
        ws[f'A{total}'] = 'Total'
        bold(ws, f'A{total}')
        if sheet.get('compare'):
            ws[f'A{total + 2}'] = 'July minus October (Rs)'
            bold(ws, f'A{total + 2}')
        if not finished:
            return
        for r in range(FIRST, last + 1):
            formula(ws, f'D{r}', f'=C{r}-B{r}')
            formula(ws, f'E{r}', f'=IF(D{r}>$B$5,$B$5,D{r})')
            formula(ws, f'F{r}', f'=D{r}-E{r}')
            formula(ws, f'G{r}', f'=E{r}*$B$3+F{r}*$B$4')
            formula(ws, f'H{r}', f'=ROUND(G{r}+$B$6,0)')
        for col in 'DEFGH':
            formula(ws, f'{col}{total}', f'=SUM({col}{FIRST}:{col}{last})')
        if sheet.get('compare'):
            formula(ws, f'B{total + 2}', f'=H{FIRST}-H{last}')
    return build


def excel_round0(x):
    """ROUND(x, 0) for the positive numbers here: halves round up, as in Excel."""
    assert abs(x - math.floor(x) - 0.5) > 1e-6, 'avoid exact halves in practice data'
    return math.floor(x + 0.5)


def expected_for(sheet, rates):
    low, high, limit, fixed = rates
    exp = {}
    rows = sheet['rows']
    for r, (_, old, new) in enumerate(rows, start=FIRST):
        units = new - old
        low_units = limit if units > limit else units
        high_units = units - low_units
        cost = low_units * low + high_units * high
        for col, v in zip('DEFGH', (units, low_units, high_units, cost,
                                    excel_round0(cost + fixed))):
            exp[f'{col}{r}'] = v
    last = FIRST + len(rows) - 1
    for col in 'DEFGH':
        exp[f'{col}{last + 1}'] = sum(exp[f'{col}{r}'] for r in range(FIRST, last + 1))
    if sheet.get('compare'):
        exp[f'B{last + 3}'] = exp[f'H{FIRST}'] - exp[f'H{last}']
    return exp


MONTHS_FINAL = RATES[:1] + (NEW_HIGH_RATE,) + RATES[2:]
flats = expected_for(FLATS, RATES)
months_before = expected_for(MONTHS, RATES)
months = expected_for(MONTHS, MONTHS_FINAL)

FILES = [
    XlsxFile('flat-bills', make_builder(FLATS, RATES), flats),
    XlsxFile('house-bills', make_builder(MONTHS, MONTHS_FINAL), months),
]


# The lesson states these results; keep them in step with the page.
def column(exp, col, first, last):
    return [round(exp[f'{col}{r}'], 6) for r in range(first, last + 1)]


assert column(flats, 'D', 9, 13) == [87, 148, 100, 213, 61]
assert column(flats, 'E', 9, 13) == [87, 100, 100, 100, 61]
assert column(flats, 'F', 9, 13) == [0, 48, 0, 113, 0]
assert column(flats, 'G', 9, 13) == [1600.8, 3179.2, 1840, 4992.7, 1122.4]
assert column(flats, 'H', 9, 13) == [1751, 3329, 1990, 5143, 1272]
assert column(flats, 'D', 14, 14) + [round(flats[f'{c}14'], 6) for c in 'EFGH'] == [
    609, 448, 161, 12735.1, 13485]

assert column(months_before, 'D', 9, 12) == [266, 232, 179, 94]
assert column(months_before, 'F', 9, 12) == [166, 132, 79, 0]
assert column(months_before, 'G', 9, 12) == [6471.4, 5522.8, 4044.1, 1729.6]
assert column(months_before, 'H', 9, 12) == [6621, 5673, 4194, 1880]
assert months_before['H13'] == 18368 and months_before['B15'] == 4741
assert column(months, 'E', 9, 12) == [100, 100, 100, 94]
assert column(months, 'G', 9, 12) == [6919.6, 5879.2, 4257.4, 1729.6]
assert column(months, 'H', 9, 12) == [7070, 6029, 4407, 1880]
assert [round(months[f'{c}13'], 6) for c in 'DEFGH'] == [771, 394, 377, 18785.8, 19386]
assert months['B15'] == 5190
# Each month's old reading is the month before's new reading.
assert all(MONTHS['rows'][i][1] == MONTHS['rows'][i - 1][2] for i in range(1, 4))
# "Flat 4 used about two and a half times the units of Flat 1, but pays almost three times as much."
assert 2.4 < flats['D12'] / flats['D9'] < 2.5 and 2.9 < flats['H12'] / flats['H9'] < 3
# Check Question 3: a flat that used 120 units.
assert excel_round0(100 * RATES[0] + 20 * RATES[1] + RATES[3]) == 2548
